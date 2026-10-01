"""Multi-format report export: DOCX (existing), PDF (existing), XLSX (new).

The SIH deck promises "Report in Existing Format (PDF, Excel, DOC)" - this
module adds the Excel leg with one sheet per block (metrics, lineage,
description), keeping every figure deterministic from report_metric rows.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.config import settings
from app.db.models import Report
from app.services.export import (
    _collect_rows,
    _header_block,
    _lineage_rows,
)

HEADER_FILL = PatternFill("solid", fgColor="1A1A2E")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _sheet_from_rows(ws, rows: list[tuple[str, ...]], headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in rows:
        ws.append(list(row))
    for i, width in enumerate([34, 30, 22, 18], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def export_report_xlsx(db: Session, report: Report) -> Path:
    additive, derived = _collect_rows(report)
    lineage = _lineage_rows(db, report)

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    _sheet_from_rows(ws, [(k, v) for k, v in _header_block(report)],
                     ["Field", "Value"])

    ws2 = wb.create_sheet("Metrics")
    _sheet_from_rows(ws2, additive, ["Metric", "Value"])
    if derived:
        start = ws2.max_row + 2
        ws2.cell(row=start, column=1, value="Derived").font = Font(bold=True)
        for j, (label, value) in enumerate(derived):
            ws2.cell(row=start + 1 + j, column=1, value=label)
            ws2.cell(row=start + 1 + j, column=2, value=value)

    ws3 = wb.create_sheet("Lineage")
    _sheet_from_rows(ws3, lineage, ["Source", "Contributions", "Edge"])

    ws4 = wb.create_sheet("Description")
    desc = report.description
    ws4.append(["Section", "Text"])
    for cell in ws4[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws4.append(["Recorded facts", desc.text_recorded if desc else ""])
    ws4.append(["Interpretation", desc.text_interpretation if desc else ""])
    ws4.append(["Source", desc.source if desc else "NONE"])
    ws4.append(["Human reviewed", "YES" if desc and desc.human_reviewed else "NO"])
    ws4.column_dimensions["A"].width = 22
    ws4.column_dimensions["B"].width = 110
    ws4.cell(row=2, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws4.cell(row=3, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    out_dir = settings.exports_path
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / f"report_{report.id}_{report.level.lower()}_{report.period_start}.xlsx"
    wb.save(out_path)
    return out_path
