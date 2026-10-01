from __future__ import annotations

import logging
from datetime import datetime, timezone
from pathlib import Path

from sqlalchemy.orm import Session

from app.config import settings
from app.db.models import Report
from app.services.factsheet import build_fact_sheet

logger = logging.getLogger("minetrace.export")

METRIC_LABELS = {
    "production_milli_t": "Production (t)",
    "ob_milli_m3": "Overburden (m3)",
    "planned_min": "Planned time (min)",
    "operating_min": "Operating time (min)",
    "stoppage_min": "Stoppage (min)",
    "workers_present": "Workers present",
}
DERIVED_LABELS = {
    "stripping_ratio_x1000": "Stripping ratio",
    "availability_bp": "Availability",
}


def _fmt_metric(code: str, value: int) -> str:
    if code.endswith("milli_t") or code.endswith("milli_m3"):
        return f"{value / 1000:,.3f}"
    if code.endswith("_bp"):
        return f"{value / 100:.2f}%"
    if code.endswith("x1000"):
        return f"{value / 1000:.2f}"
    return f"{value:,}"


def _collect_rows(report: Report) -> tuple[list[tuple[str, str]], list[tuple[str, str]]]:
    additive: list[tuple[str, str]] = []
    derived: list[tuple[str, str]] = []
    for m in sorted(report.metrics, key=lambda x: x.metric_code):
        label = METRIC_LABELS.get(m.metric_code, m.metric_code)
        if m.metric_code in DERIVED_LABELS:
            label = DERIVED_LABELS[m.metric_code]
        row = (label, _fmt_metric(m.metric_code, m.value_int))
        if m.is_additive:
            additive.append(row)
        else:
            derived.append(row)
    return additive, derived


def _lineage_rows(db: Session, report: Report) -> list[tuple[str, str, str]]:
    rows: list[tuple[str, str, str]] = []
    for edge in report.lineage_edges:
        if edge.child_kind == "SHIFT_ENTRY":
            child = f"Shift entry #{edge.child_shift_entry_id}"
        else:
            child = f"Report #{edge.child_report_id}"
        contribs: list[str] = []
        for mc in edge.contributions:
            metric = mc.parent_metric
            label = METRIC_LABELS.get(metric.metric_code, metric.metric_code)
            value = _fmt_metric(metric.metric_code, mc.contributed_value_int)
            contribs.append(f"{label}: {value}")
        rows.append((child, ", ".join(contribs) or "-", f"edge #{edge.id}"))
    return rows


def _header_block(report: Report) -> list[tuple[str, str]]:
    period = str(report.period_start)
    if report.period_start != report.period_end:
        period = f"{report.period_start} to {report.period_end}"
    return [
        ("Report level", report.level),
        ("Mine ID", str(report.mine_id)),
        ("Bench ID", str(report.bench_id) if report.bench_id is not None else "All benches"),
        ("Period", period),
        ("Version", f"v{report.version} ({report.status})"),
        ("Data coverage", f"{report.present_children}/{report.expected_children} children"),
        ("Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
    ]


def export_report(db: Session, report: Report, fmt: str) -> Path:
    fmt = fmt.lower()
    if fmt not in ("docx", "pdf"):
        raise ValueError(f"Unsupported format: {fmt}")

    facts = build_fact_sheet(db, report)
    additive, derived = _collect_rows(report)
    lineage = _lineage_rows(db, report)

    description = report.description
    recorded = description.text_recorded if description else ""
    interpretation = description.text_interpretation if description else ""
    desc_source = description.source if description else "NONE"

    stem = f"report_{report.id}_{report.level.lower()}_{report.period_start}"
    out_dir = settings.exports_path
    out_dir.mkdir(parents=True, exist_ok=True)

    if fmt == "docx":
        out_path = out_dir / f"{stem}.docx"
        _write_docx(
            out_path, report, facts, additive, derived,
            lineage, recorded, interpretation, desc_source,
        )
    else:
        out_path = out_dir / f"{stem}.pdf"
        _write_pdf(
            out_path, report, facts, additive, derived,
            lineage, recorded, interpretation, desc_source,
        )

    return out_path


def _write_docx(
    out_path: Path,
    report: Report,
    facts: dict,
    additive: list[tuple[str, str]],
    derived: list[tuple[str, str]],
    lineage: list[tuple[str, str, str]],
    recorded: str,
    interpretation: str,
    desc_source: str,
) -> None:
    from docx import Document as DocxDocument
    from docx.shared import Pt

    doc = DocxDocument()
    doc.core_properties.title = f"MineTrace {report.level} Report {report.period_start}"

    title = doc.add_heading(f"MineTrace {report.level} Report", level=0)
    title.alignment = 0

    doc.add_heading("Report Information", level=1)
    for label, value in _header_block(report):
        p = doc.add_paragraph()
        run = p.add_run(f"{label}: ")
        run.bold = True
        p.add_run(value)

    doc.add_heading("Additive Metrics", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text = "Metric", "Value"
    for label, value in additive:
        row = table.add_row().cells
        row[0].text, row[1].text = label, value

    if derived:
        doc.add_heading("Derived Metrics", level=1)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text = "Metric", "Value"
        for label, value in derived:
            row = table.add_row().cells
            row[0].text, row[1].text = label, value

    doc.add_heading("Description", level=1)
    if recorded:
        doc.add_heading("Recorded facts", level=2)
        doc.add_paragraph(recorded)
    if interpretation:
        doc.add_heading("Interpretation", level=2)
        doc.add_paragraph(interpretation)
    if not recorded and not interpretation:
        doc.add_paragraph("No narration generated yet.")
    doc.add_paragraph(f"Narration source: {desc_source}").runs[0].italic = True

    if lineage:
        doc.add_heading("Lineage — Source Contributions", level=1)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Source", "Contributions", "Edge"
        for child, contribs, edge in lineage:
            row = table.add_row().cells
            row[0].text, row[1].text, row[2].text = child, contribs, edge

    doc.add_heading("Traceability Statement", level=1)
    para = doc.add_paragraph(
        "Every figure in this report is computed deterministically from APPROVED "
        "shift entries. The full contribution trail is recorded in the lineage "
        "tables above; each entry can be traced to its source record by ID. "
        "Narrative text was validated against the fact sheet by NumberGuard; "
        "no AI-generated number appears in this report."
    )
    para.runs[0].font.size = Pt(9)

    doc.save(out_path)


def _write_pdf(
    out_path: Path,
    report: Report,
    facts: dict,
    additive: list[tuple[str, str]],
    derived: list[tuple[str, str]],
    lineage: list[tuple[str, str, str]],
    recorded: str,
    interpretation: str,
    desc_source: str,
) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleCentered", parent=styles["Title"], alignment=TA_CENTER,
    )
    small = ParagraphStyle("Small", parent=styles["Normal"], fontSize=8, leading=10)

    story: list = [
        Paragraph(f"MineTrace {report.level} Report", title_style),
        Spacer(1, 6),
        Paragraph(f"Period: {dict(_header_block(report))['Period']}", styles["Heading2"]),
        Spacer(1, 8),
    ]

    def kv_table(rows: list[tuple[str, str]]) -> Table:
        data = [[Paragraph("<b>Metric</b>", styles["Normal"]),
                 Paragraph("<b>Value</b>", styles["Normal"])]]
        for label, value in rows:
            data.append([Paragraph(label, styles["Normal"]),
                         Paragraph(value, styles["Normal"])])
        t = Table(data, colWidths=[280, 220])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a1a2e")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f6fa")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ]))
        return t

    story.append(Paragraph("Additive Metrics", styles["Heading2"]))
    story.append(kv_table(additive))
    if derived:
        story.append(Paragraph("Derived Metrics", styles["Heading2"]))
        story.append(kv_table(derived))

    story.append(Paragraph("Description", styles["Heading2"]))
    if recorded:
        story.append(Paragraph("<b>Recorded facts</b>", styles["Normal"]))
        story.append(Paragraph(recorded.replace("\n", "<br/>"), small))
    if interpretation:
        story.append(Paragraph("<b>Interpretation</b>", styles["Normal"]))
        story.append(Paragraph(interpretation.replace("\n", "<br/>"), small))
    if not recorded and not interpretation:
        story.append(Paragraph("No narration generated yet.", small))
    story.append(Paragraph(f"Narration source: {desc_source}", small))

    if lineage:
        story.append(Paragraph("Lineage — Source Contributions", styles["Heading2"]))
        data = [[Paragraph("<b>Source</b>", styles["Normal"]),
                 Paragraph("<b>Contributions</b>", styles["Normal"]),
                 Paragraph("<b>Edge</b>", styles["Normal"])]]
        for child, contribs, edge in lineage:
            data.append([
                Paragraph(child, small),
                Paragraph(contribs, small),
                Paragraph(edge, small),
            ])
        t = Table(data, colWidths=[130, 290, 80], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a1a2e")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ]))
        story.append(t)

    story.append(Spacer(1, 10))
    story.append(Paragraph(
        "Traceability: every figure is computed deterministically from APPROVED shift "
        "entries and is traceable to its source record via the lineage edges above. "
        "Narrative text was validated by NumberGuard; no AI-generated number appears "
        "in this report.",
        small,
    ))

    doc = SimpleDocTemplate(
        str(out_path), pagesize=A4,
        leftMargin=42, rightMargin=42, topMargin=42, bottomMargin=42,
    )
    doc.build(story)
