from __future__ import annotations

import logging
from pathlib import Path

from sqlalchemy.orm import Session

from app.config import settings
from app.db.models import Document
from app.services import storage

logger = logging.getLogger("minetrace.ingestion")

ALLOWED_TYPES = {
    ".pdf": "PDF",
    ".png": "IMAGE",
    ".jpg": "IMAGE",
    ".jpeg": "IMAGE",
    ".tif": "IMAGE",
    ".tiff": "IMAGE",
    ".xlsx": "SPREADSHEET",
    ".xls": "SPREADSHEET",
    ".csv": "SPREADSHEET",
}


def _read_spreadsheet_rows(path: Path) -> list[list[str]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        import csv

        with open(path, newline="", encoding="utf-8", errors="replace") as f:
            return [row for row in csv.reader(f)]

    try:
        import openpyxl
    except ImportError:
        logger.warning("openpyxl not installed; cannot read %s", path.name)
        return []

    rows: list[list[str]] = []
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows.append([f"=== Sheet: {ws.title} ==="])
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v) for v in row]
                if any(c.strip() for c in cells):
                    rows.append(cells)
        wb.close()
    except Exception:
        logger.exception("Failed to parse spreadsheet %s", path.name)
    return rows


def save_upload(
    db: Session, filename: str, content: bytes, user_id: int,
) -> Document:
    suffix = Path(filename).suffix.lower()
    file_type = ALLOWED_TYPES.get(suffix)
    if not file_type:
        msg = f"Unsupported file type: {suffix}"
        raise ValueError(msg)

    key = storage.new_key(filename)
    storage.save_bytes(key, content)

    doc = Document(
        filename=filename,
        file_path=key,
        file_type=file_type,
        file_size=len(content),
        uploaded_by=user_id,
        status="UPLOADED",
    )
    db.add(doc)
    db.flush()
    return doc


def extract_text(doc: Document) -> list[dict]:
    """Extract page text from the local filesystem or the object store
    (S3/MinIO bytes are materialized to a temp file for the parsers)."""
    raw = Path(doc.file_path)
    temp: Path | None = None
    if raw.exists():
        path = raw
    else:
        content = storage.open_bytes(doc.file_path)
        temp = settings.page_cache_path / f"extract_{doc.id}{raw.suffix.lower()}"
        temp.write_bytes(content)
        path = temp
    try:
        if doc.file_type == "PDF":
            return _extract_pdf(path)
        if doc.file_type == "SPREADSHEET":
            return _extract_spreadsheet(path)
        return _extract_image(path)
    finally:
        if temp is not None:
            temp.unlink(missing_ok=True)


def _extract_spreadsheet(path: Path) -> list[dict]:
    """Flatten spreadsheet rows into page-like text blocks so they flow
    through the same chunk -> embed -> RAG pipeline as PDFs and images."""
    rows = _read_spreadsheet_rows(path)
    if not rows:
        return [{"page": 1, "text": ""}]

    # Header propagation: repeat the first meaningful row on every page so
    # each chunk is self-describing (retrieval + LLM can interpret values).
    header: str | None = None
    body_lines: list[str] = []
    for row in rows:
        line = " | ".join(cell.strip() for cell in row).strip(" |")
        if not line:
            continue
        if header is None and not line.startswith("==="):
            header = line
            continue
        if line.startswith("==="):
            continue  # sheet markers: not corpus-relevant
        body_lines.append(line)

    pages: list[dict] = []
    current_lines: list[str] = []
    target_chars = 3000

    for line in body_lines:
        current_lines.append(line)
        if sum(len(item) for item in current_lines) >= target_chars:
            pages.append({"page": len(pages) + 1,
                          "text": (f"{header}\n" if header else "") + "\n".join(current_lines)})
            current_lines = []
    if current_lines:
        pages.append({"page": len(pages) + 1,
                      "text": (f"{header}\n" if header else "") + "\n".join(current_lines)})

    return pages or [{"page": 1, "text": header or ""}]


def _extract_pdf(path: Path) -> list[dict]:
    pages = []
    try:
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            for i, page in enumerate(pdf.pages):
                text = page.extract_text() or ""
                if not text.strip():
                    text = _ocr_page_image(path, i)
                pages.append({"page": i + 1, "text": text})
    except ImportError:
        logger.warning("pdfplumber not installed, using OCR only")
        pages = [{"page": 1, "text": _ocr_file(path)}]
    return pages


def _extract_image(path: Path) -> list[dict]:
    text = _ocr_file(path)
    return [{"page": 1, "text": text}]


_ocr_engine = None


def _get_ocr():
    global _ocr_engine
    if _ocr_engine is None:
        from paddleocr import PaddleOCR

        logger.info("Initializing PaddleOCR (once)")
        _ocr_engine = PaddleOCR(use_angle_cls=True, lang="en", show_log=False)
    return _ocr_engine


def _ocr_file(path: Path) -> str:
    text = _ocr_paddle(path)
    if text:
        return text
    return _ocr_tesseract(path)


def _ocr_paddle(path: Path) -> str:
    try:
        ocr = _get_ocr()
        result = ocr.ocr(str(path), cls=True)
        lines = []
        if result and result[0]:
            for line in result[0]:
                if line and len(line) >= 2:
                    lines.append(line[1][0])
        return "\n".join(lines)
    except ImportError:
        logger.warning("PaddleOCR not installed")
        return ""
    except Exception:
        logger.exception("OCR failed for %s", path)
        return ""


def _ocr_tesseract(path: Path) -> str:
    """Fallback OCR via pytesseract (activates when the tesseract binary is
    installed and PaddleOCR is unavailable)."""
    try:
        import pytesseract
        from PIL import Image

        return pytesseract.image_to_string(Image.open(path)) or ""
    except Exception:
        return ""


def _ocr_page_image(pdf_path: Path, page_idx: int) -> str:
    try:
        import pdfplumber

        with pdfplumber.open(pdf_path) as pdf:
            page = pdf.pages[page_idx]
            img = page.to_image(resolution=300)
            tmp = settings.page_cache_path / f"ocr_{page_idx}.png"
            img.save(tmp)
            text = _ocr_file(tmp)
            tmp.unlink(missing_ok=True)
            return text
    except Exception:
        logger.exception("Page OCR failed")
        return ""


def delete_document(db: Session, doc: Document) -> None:
    raw = Path(doc.file_path)
    if raw.exists():
        raw.unlink(missing_ok=True)
    else:
        try:
            storage.delete(doc.file_path)
        except Exception:
            logger.warning("could not delete stored object %s", doc.file_path)
    db.delete(doc)
